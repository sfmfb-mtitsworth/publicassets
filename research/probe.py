"""One-off probe: find EDD's weekly UI claims XLSX files and show their layout for SF and Marin."""
import io, re, sys, urllib.parse
sys.path.insert(0, "etl")
from common import http_get
import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (compatible; sfmfb-leadingindicators)"}
pages = ["https://edd.ca.gov/en/newsroom/facts-and-stats/dashboard/", "https://edd.ca.gov/Newsroom/facts-and-stats/dashboard.htm",
         "https://edd.ca.gov/en/about_edd/Quick_Statistics_Information_by_County"]
links = set()
for pg in pages:
    try:
        html = http_get(pg, headers=UA, retries=1).decode("utf-8", "replace")
        found = re.findall(r'href="([^"]+\.(?:xlsx?|csv))"', html, re.I)
        print(f"== {pg}: {len(found)} spreadsheet links")
        for l in found:
            u = urllib.parse.urljoin(pg, l); links.add(u); print("   ", u)
    except Exception as e:
        print(f"== {pg}: {e}")
for u in sorted(links):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(http_get(u, headers=UA, retries=1)), read_only=True, data_only=True)
    except Exception as e:
        print(f"\n## {u}: {e}"); continue
    print(f"\n## {u}\n   sheets: {wb.sheetnames}")
    for ws in wb.worksheets[:6]:
        rows = list(ws.iter_rows(max_row=4000, values_only=True))
        print(f"   -- sheet {ws.title!r}: {len(rows)} rows")
        for i, r in enumerate(rows[:8]):
            print(f"      r{i}: {[str(c)[:28] for c in r[:14] if c is not None]}")
        hits = [r for r in rows if any(str(c).strip().lower() in ("san francisco", "marin", "san francisco county", "marin county") for c in r if c is not None)]
        for r in hits[:4] + hits[-4:]:
            print(f"      SF/Marin: {[str(c)[:20] for c in r[:14] if c is not None]}")
        if hits: print(f"      ({len(hits)} SF/Marin rows)")
